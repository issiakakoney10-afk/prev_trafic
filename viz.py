"""
Extraction des données déjà calculées d'un classeur (valeurs, pas formules)
en tableaux pandas, pour l'onglet Visualisation de l'application. Ce module
ne dépend pas de Streamlit : il ne fait que lire un classeur et rendre des
DataFrames simples, réutilisables ailleurs si besoin.
"""
from pathlib import Path

import openpyxl
import pandas as pd

from common import (
    ANCHOR_ROWS, ANNOUNCED_MONTH_COLS, ANNUEL_REALISE_ROWS, CALC_PREV_MONTH_COLS,
    CALC_PREV_REALISE_ROWS, CALC_PREV_TRAFIC_MENS_ROWS, DIRECT_MONTH_COLS,
    MONTHS_FR, OBJECTIFS_ANNUELS, OBJECTIFS_SHEET, TC2_ENGAGEMENT_ROW,
    TC2_SHEET, TC2_YEARS_COLS, YEAR_SHEET,
)

CALC_PREV_SHEET = "CALC. PREV"
ANNUEL_SHEET = "ANNUEL"

KPI_LABELS = {
    "escales": "Escales de navires",
    "tonnage_global": "Trafic global (tonnes)",
    "national": "Trafic national (tonnes)",
    "transit": "Trafic transit (tonnes)",
    "transbordement": "Trafic transbordement (tonnes)",
    "conteneurs_teu": "Conteneurs (EVP)",
    "conteneurs_transbordes_teu": "Conteneurs transbordés (EVP)",
}
# ordre fixe : chaque KPI garde toujours la même couleur dans toute l'appli
KPI_ORDER = list(CALC_PREV_REALISE_ROWS.keys())


def load_wb(path, data_only=True):
    return openpyxl.load_workbook(Path(path), data_only=data_only)


def _month_series(ws, row, cols=CALC_PREV_MONTH_COLS):
    return [ws[f"{c}{row}"].value or 0 for c in cols]


def calc_prev_df(wb) -> pd.DataFrame:
    """Table longue : KPI x mois x valeur, ligne 'Réalisé' de CALC. PREV."""
    ws = wb[CALC_PREV_SHEET]
    rows = []
    for kpi, r in CALC_PREV_REALISE_ROWS.items():
        vals = _month_series(ws, r)
        for m, v in zip(MONTHS_FR, vals):
            rows.append({"kpi": kpi, "KPI": KPI_LABELS[kpi], "Mois": m, "Valeur": v})
    return pd.DataFrame(rows)


def annuel_df(wb_values, wb_formulas) -> pd.DataFrame:
    """Table longue : KPI x mois x valeur pour ANNUEL, + statut
    (Réalisé / Projeté) déduit des formules elles-mêmes (pas d'hypothèse
    extérieure) : si la formule pointe vers la ligne 'Réalisé' de CALC. PREV
    -> réalisé ; si elle pointe vers ligne+decalage -> projeté (méthode).

    Attention : la ligne à LIRE dans ANNUEL (ANNUEL_REALISE_ROWS) n'est pas
    la même que la ligne 'Réalisé' de CALC. PREV (CALC_PREV_REALISE_ROWS) —
    voir le commentaire dans common.py. On lit la première, on compare la
    formule trouvée à la seconde pour déduire le statut."""
    wsv = wb_values[ANNUEL_SHEET]
    wsf = wb_formulas[ANNUEL_SHEET]
    rows = []
    for kpi, calc_prev_realise_row in CALC_PREV_REALISE_ROWS.items():
        annuel_row = ANNUEL_REALISE_ROWS[kpi]
        for col, mois in zip(CALC_PREV_MONTH_COLS, MONTHS_FR):
            val = wsv[f"{col}{annuel_row}"].value or 0
            formula = wsf[f"{col}{annuel_row}"].value or ""
            statut = "Réalisé"
            if isinstance(formula, str):
                digits = "".join(ch for ch in formula if ch.isdigit())
                if digits and int(digits) != calc_prev_realise_row:
                    statut = "Projeté"
            rows.append({"kpi": kpi, "KPI": KPI_LABELS[kpi], "Mois": mois,
                         "Valeur": val, "Statut": statut})
    return pd.DataFrame(rows)


def trafic_mens_kpi_df(wb, sheet_name: str) -> pd.DataFrame:
    """Table longue : KPI x mois pour une feuille 'Trafic mensXXXX' (lignes
    de détail direct, colonnes C:N)."""
    if sheet_name not in wb.sheetnames:
        return pd.DataFrame(columns=["kpi", "KPI", "Mois", "Valeur"])
    ws = wb[sheet_name]
    rows = []
    for kpi, refs in CALC_PREV_TRAFIC_MENS_ROWS.items():
        vals = _month_series(ws, refs["direct"], DIRECT_MONTH_COLS)
        for m, v in zip(MONTHS_FR, vals):
            rows.append({"kpi": kpi, "KPI": KPI_LABELS[kpi], "Mois": m, "Valeur": v})
    return pd.DataFrame(rows)


def seasonal_block(wb, cell: str):
    """Un bloc de COEFFICIENTS_SAISONNIERS_2026 (ex. cell='J4') : renvoie
    (historique 5 ans en DataFrame long [Année, Mois, Valeur],
     coefficients mensuels en DataFrame [Mois, Coef.],
     objectif annuel (float))."""
    ws = wb[OBJECTIFS_SHEET]
    header_row = int("".join(ch for ch in cell if ch.isdigit()))
    objectif = ws[cell].value
    hist_rows = []
    for offset in range(2, 7):  # 5 années d'historique
        year = ws[f"B{header_row + offset}"].value
        vals = _month_series(ws, header_row + offset, DIRECT_MONTH_COLS)
        for m, v in zip(MONTHS_FR, vals):
            hist_rows.append({"Année": year, "Mois": m, "Valeur": v})
    coef_row = header_row + 9
    coefs = _month_series(ws, coef_row, DIRECT_MONTH_COLS)
    coef_df = pd.DataFrame({"Mois": MONTHS_FR, "Coefficient": coefs})
    return pd.DataFrame(hist_rows), coef_df, objectif


def tc2_df(wb) -> pd.DataFrame:
    """Engagement contractuel TC2 vs trafic transbordement conteneurisé
    projeté (business plan), par année."""
    ws = wb[TC2_SHEET]
    rows = []
    for year, col in sorted(TC2_YEARS_COLS.items()):
        engagement = ws[f"{col}{TC2_ENGAGEMENT_ROW}"].value
        plan = ws[f"{col}39"].value  # ligne "TRANSBO" du business plan
        rows.append({"Année": year, "Engagement contractuel": engagement,
                     "Plan d'affaires (projeté)": plan})
    return pd.DataFrame(rows)


def repartition_df(wb) -> pd.DataFrame:
    """Quelques séries annuelles clés de REPARTITION (2019 -> 2026)."""
    ws = wb["REPARTITION"]
    years = [ws.cell(row=3, column=c).value for c in range(2, 10)]
    def series(row, label):
        vals = [ws.cell(row=row, column=c).value for c in range(2, 10)]
        return [{"Année": y, "Série": label, "Valeur": v}
                for y, v in zip(years, vals) if y and v is not None]
    rows = []
    rows += series(4, "Escales navires (nombre)")
    rows += series(5, "  dont Port de Commerce")
    rows += series(6, "  dont Port de Pêche")
    rows += series(16, "Trafic global (tonnes)")
    rows += series(30, "Conteneurs (EVP)")
    return pd.DataFrame(rows)


STATS_EXO_SECTIONS = {
    "Import / Export (tonnes)": {"header": 5, "rows": [(8, "Importations"), (20, "Exportations")]},
    "Port de Commerce / Pêche (tonnes)": {"header": 47, "rows": [(50, "Port de Commerce"), (55, "Terminal à Pêche")]},
    "Tare des conteneurs (tonnes)": {"header": 5, "rows": [(10, "  National"), (11, "  Transbordement"), (12, "  Transit")]},
}


def stats_exo_df(wb, section_key: str) -> pd.DataFrame:
    ws = wb["Stats exo 2025brute"]
    section = STATS_EXO_SECTIONS[section_key]
    header_row = section["header"]
    years = [ws.cell(row=header_row + 1, column=c).value for c in range(3, 7)]
    rows = []
    for row, label in section["rows"]:
        for y, col in zip(years, range(3, 7)):
            v = ws.cell(row=row, column=col).value
            if y is not None and v is not None:
                rows.append({"Année": int(y), "Catégorie": label, "Valeur": v})
    return pd.DataFrame(rows)
